#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""
Generator za ROYAL_NANNY_Obuka_Predavaci.pptx
Kreira luksuznu, profesionalnu 16:9 prezentaciju sa ugrađenim beleškama za predavača.
"""

import os
import re
import pptx
from pptx.util import Inches, Pt
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

from slides_data import SLIDES_DATA

BASE_DIR = "/Users/markozivkovic/PREZENTACIJA ZA ROYAL NANNY"
ASSETS_DIR = os.path.join(BASE_DIR, "extracted_assets")
OUTPUT_PPTX = os.path.join(BASE_DIR, "ROYAL_NANNY_Obuka_Predavaci.pptx")

# Boje
RGB_DARK = RGBColor(71, 52, 46)        # #47342E - Tamna espresso
RGB_BG = RGBColor(247, 245, 240)       # #F7F5F0 - Topla krem
RGB_SAND = RGBColor(205, 187, 160)     # #CDBBA0 - Pesak bež
RGB_LINEN = RGBColor(240, 233, 215)    # #F0E9D7 - Svetli lan
RGB_TERRA = RGBColor(118, 85, 53)      # #765535 - Konjak / terakota
RGB_SAGE = RGBColor(139, 144, 109)     # #8B906D - Žalfija
RGB_OLIVE = RGBColor(105, 97, 79)      # #69614F - Maslinasto siva
RGB_GOLD = RGBColor(209, 174, 102)     # #D1AE66 - Oker zlato
RGB_BLUSH = RGBColor(231, 219, 206)    # #E7DBCE - Nežna rumen
RGB_WHITE = RGBColor(255, 255, 255)
RGB_BORDER = RGBColor(225, 218, 206)
RGB_CARD_BG = RGBColor(255, 255, 255)
RGB_DARK_CARD = RGBColor(58, 43, 38)

FONT_TITLE = "Libre Baskerville"
FONT_BODY = "Poppins"

def clean_html(text):
    """Uklanja HTML tagove za PPTX plain text format."""
    text = re.sub(r'</?(?:strong|b)>', '', text)
    text = re.sub(r'</?(?:em|i)>', '', text)
    text = re.sub(r'<br\s*/?>', '\n', text)
    return text

def strip_em_for_plain(text):
    return text.replace('<em>', '').replace('</em>', '')

def set_slide_background(slide, color):
    background = slide.background
    fill = background.fill
    fill.solid()
    fill.fore_color.rgb = color

def add_header(slide, s_data, is_dark=False):
    """Kreira elegantno zaglavlje sa monogramom, modulom, kategorijom i logotipom."""
    # Monogram ikonica levo (kao na stranici 12 brend knjige)
    monogram_file = "logo_mark_cream.png" if is_dark else "logo_mark_transparent.png"
    monogram_path = os.path.join(ASSETS_DIR, monogram_file)
    if os.path.exists(monogram_path):
        slide.shapes.add_picture(monogram_path, Inches(0.8), Inches(0.32), width=Inches(0.38), height=Inches(0.42))
        
    # Kategorija / Modul badge
    txBox = slide.shapes.add_textbox(Inches(1.28), Inches(0.36), Inches(8.0), Inches(0.35))
    tf = txBox.text_frame
    tf.margin_left = tf.margin_top = tf.margin_right = tf.margin_bottom = 0
    p = tf.paragraphs[0]
    p.text = f"{s_data.get('category', 'ROYAL NANNY')} • {s_data.get('module', '')}".upper()
    p.font.name = FONT_BODY
    p.font.size = Pt(10)
    p.font.bold = True
    p.font.color.rgb = RGB_SAND if is_dark else RGB_TERRA
    
    # Logo u gornjem desnom uglu
    logo_file = "logo_horizontal_light.png" if is_dark else "logo_horizontal_transparent.png"
    logo_path = os.path.join(ASSETS_DIR, logo_file)
    if os.path.exists(logo_path):
        slide.shapes.add_picture(logo_path, Inches(10.6), Inches(0.32), width=Inches(1.9))

def add_footer(slide, s_data, current_idx, total_slides, is_dark=False):
    """Kreira diskretno podnožje sa brend vrednostima, modulom i brojem slajda."""
    # Leva strana - brend vrednosti (kao na stranici 12 brend knjige)
    txValues = slide.shapes.add_textbox(Inches(0.8), Inches(7.05), Inches(5.5), Inches(0.3))
    tfValues = txValues.text_frame
    tfValues.margin_left = tfValues.margin_top = tfValues.margin_right = tfValues.margin_bottom = 0
    pVal = tfValues.paragraphs[0]
    pVal.text = "PROVERENE   •   ELEGANTNE   •   POUZDANE"
    pVal.font.name = FONT_BODY
    pVal.font.size = Pt(9)
    pVal.font.bold = True
    pVal.font.color.rgb = RGB_SAND if is_dark else RGB_TERRA

    # Desna strana - Akademija & Broj slajda
    txRight = slide.shapes.add_textbox(Inches(6.5), Inches(7.05), Inches(6.0), Inches(0.3))
    tfRight = txRight.text_frame
    tfRight.margin_left = tfRight.margin_top = tfRight.margin_right = tfRight.margin_bottom = 0
    pRight = tfRight.paragraphs[0]
    pRight.text = f"ROYAL NANNY AKADEMIJA • STANDARDI NEGE   |   {current_idx:02d} / {total_slides:02d}"
    pRight.alignment = PP_ALIGN.RIGHT
    pRight.font.name = FONT_BODY
    pRight.font.size = Pt(9)
    pRight.font.color.rgb = RGB_SAND if is_dark else RGB_OLIVE

def add_notes(slide, s_data):
    """Dodaje detaljne beleške za predavača u Speaker Notes polje."""
    notes = s_data.get("lecturer_notes", {})
    if not notes:
        return
    notes_slide = slide.notes_slide
    tf = notes_slide.notes_text_frame
    
    notes_text = (
        f"SLAJD {s_data['id']}: {s_data['title']}\n"
        f"Modul: {s_data.get('module', '')} | Kategorija: {s_data.get('category', '')}\n\n"
        f"🎯 CILJ OVE TEME:\n{notes.get('cilj', '')}\n\n"
        f"📌 KLJUČNE TEZE I SMERNICE ZA PREDAVAČA:\n{notes.get('teze', '')}\n\n"
        f"⚠️ NAJČEŠĆE GREŠKE I ZABLUDE U PRAKSI:\n{notes.get('greske', '')}\n\n"
        f"❓ INTERAKTIVNO PITANJE ZA POLAZNICE:\n{notes.get('pitanje', '')}\n"
    )
    tf.text = notes_text

def build_pptx():
    print(f"Pokrećem generisanje PowerPoint prezentacije: {OUTPUT_PPTX}")
    prs = pptx.Presentation()
    prs.slide_width = Inches(13.333)
    prs.slide_height = Inches(7.5)
    blank_layout = prs.slide_layouts[6] # blank
    
    total = len(SLIDES_DATA)
    
    for idx, s in enumerate(SLIDES_DATA, 1):
        slide = prs.slides.add_slide(blank_layout)
        is_dark = s.get("bg_theme") == "dark"
        
        # 1. Pozadina
        set_slide_background(slide, RGB_DARK if is_dark else RGB_BG)
        
        # 2. Header & Footer
        add_header(slide, s, is_dark)
        add_footer(slide, s, idx, total, is_dark)
        
        # 3. Notes
        add_notes(slide, s)
        
        # 4. Naslov i podnaslov
        title_box = slide.shapes.add_textbox(Inches(0.8), Inches(0.72), Inches(7.4 if s["layout"] == "split_right_image" else 11.7), Inches(1.15))
        tf_title = title_box.text_frame
        tf_title.word_wrap = True
        tf_title.margin_left = tf_title.margin_top = tf_title.margin_right = tf_title.margin_bottom = 0
        
        p_title = tf_title.paragraphs[0]
        p_title.text = s["title"]
        p_title.font.name = FONT_TITLE
        p_title.font.size = Pt(29 if len(s["title"]) > 38 else 34)
        p_title.font.bold = True
        p_title.font.color.rgb = RGB_LINEN if is_dark else RGB_DARK
        
        if s.get("subtitle"):
            p_sub = tf_title.add_paragraph()
            p_sub.text = s["subtitle"]
            p_sub.font.name = FONT_BODY
            p_sub.font.size = Pt(13.5)
            p_sub.font.color.rgb = RGB_SAND if is_dark else RGB_OLIVE
            p_sub.space_before = Pt(3)
            
        # 5. Render po layout-u
        layout = s.get("layout", "two_col_cards")
        
        if layout == "cover":
            render_cover_layout(slide, s)
        elif layout == "split_right_image":
            render_split_image_layout(slide, s, is_dark)
        elif layout == "two_col_cards":
            render_two_col_layout(slide, s, is_dark)
        elif layout == "three_cards":
            render_three_cards_layout(slide, s, is_dark)
        elif layout == "matrix_4":
            render_matrix_layout(slide, s, is_dark)
            
    prs.save(OUTPUT_PPTX)
    print(f"✅ Uspešno sačuvan PowerPoint fajl: {OUTPUT_PPTX} ({len(prs.slides)} slajdova)")

def render_cover_layout(slide, s):
    # Desna velika slika
    img_name = s.get("image")
    if img_name:
        img_path = os.path.join(ASSETS_DIR, img_name)
        if os.path.exists(img_path):
            card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(7.8), Inches(1.85), Inches(4.7), Inches(5.0))
            card.fill.solid()
            card.fill.fore_color.rgb = RGB_DARK_CARD
            card.line.color.rgb = RGB_SAND
            card.line.width = Pt(1.5)
            slide.shapes.add_picture(img_path, Inches(7.95), Inches(2.0), width=Inches(4.4), height=Inches(4.7))
            
    # Leva strana - opis
    card_l = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), Inches(1.85), Inches(6.7), Inches(2.8))
    card_l.fill.solid()
    card_l.fill.fore_color.rgb = RGB_DARK_CARD
    card_l.line.color.rgb = RGB_SAND
    card_l.line.width = Pt(1)
    
    bar_l = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), Inches(1.85), Inches(0.08), Inches(2.8))
    bar_l.fill.solid()
    bar_l.fill.fore_color.rgb = RGB_GOLD
    bar_l.line.fill.background()
    
    tf = card_l.text_frame
    tf.word_wrap = True
    tf.margin_left = Inches(0.35)
    tf.margin_right = Inches(0.3)
    tf.margin_top = Inches(0.25)
    tf.margin_bottom = Inches(0.25)
    
    blocks = s.get("content_blocks", [])
    if blocks:
        b = blocks[0]
        p1 = tf.paragraphs[0]
        p1.text = b.get("title", "")
        p1.font.name = FONT_TITLE
        p1.font.size = Pt(22)
        p1.font.bold = True
        p1.font.color.rgb = RGB_SAND
        
        p2 = tf.add_paragraph()
        p2.text = clean_html(b.get("text", ""))
        p2.font.name = FONT_BODY
        p2.font.size = Pt(14.5)
        p2.font.color.rgb = RGB_LINEN
        p2.space_before = Pt(8)
        
        if b.get("bullets"):
            for bullet in b["bullets"]:
                pb = tf.add_paragraph()
                pb.text = "• " + clean_html(bullet)
                pb.font.name = FONT_BODY
                pb.font.size = Pt(11.5)
                pb.font.color.rgb = RGB_LINEN
                pb.space_before = Pt(4)
                
    # Metrics kartice na dnu leve strane
    metrics = s.get("metrics", [])
    if metrics:
        left_pos = Inches(0.8)
        for m in metrics:
            m_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, Inches(4.85), Inches(2.1), Inches(2.0))
            m_card.fill.solid()
            m_card.fill.fore_color.rgb = RGB_DARK_CARD
            m_card.line.color.rgb = RGB_SAND
            m_card.line.width = Pt(1)
            
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, Inches(4.85), Inches(0.07), Inches(2.0))
            bar.fill.solid()
            bar.fill.fore_color.rgb = RGB_GOLD
            bar.line.fill.background()
            
            mtf = m_card.text_frame
            mtf.word_wrap = True
            mtf.margin_left = Inches(0.2)
            mtf.margin_right = Inches(0.15)
            mtf.margin_top = Inches(0.25)
            
            mp1 = mtf.paragraphs[0]
            mp1.text = m["val"]
            mp1.font.name = FONT_TITLE
            mp1.font.size = Pt(32)
            mp1.font.bold = True
            mp1.font.color.rgb = RGB_GOLD
            
            mp2 = mtf.add_paragraph()
            mp2.text = m["lbl"]
            mp2.font.name = FONT_BODY
            mp2.font.size = Pt(12)
            mp2.font.color.rgb = RGB_LINEN
            mp2.space_before = Pt(6)
            
            left_pos += Inches(2.25)

def render_split_image_layout(slide, s, is_dark):
    # Desna slika
    img_name = s.get("image")
    if img_name:
        img_path = os.path.join(ASSETS_DIR, img_name)
        if os.path.exists(img_path):
            is_info = img_name.startswith("slide") or any(k in img_name for k in ["_protocol", "_card", "_station", "_care", "_admin", "_kit", "_allergies", "_positions", "_first_aid", "_dehydration", "_hd.png"])
            if not is_info:
                card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(8.0), Inches(1.85), Inches(4.55), Inches(5.0))
                card.fill.solid()
                card.fill.fore_color.rgb = RGB_DARK_CARD if is_dark else RGB_CARD_BG
                card.line.color.rgb = RGB_SAND if is_dark else RGB_BORDER
                card.line.width = Pt(1)
            
            from PIL import Image as PILImage
            with PILImage.open(img_path) as p_img:
                img_w, img_h = p_img.size
                aspect = img_w / img_h
                
            box_w = Inches(4.25)
            box_h = Inches(4.7)
            box_aspect = 4.25 / 4.7
            
            if aspect > box_aspect:
                pic_w = box_w
                pic_h = box_w / aspect
                pic_left = Inches(8.15)
                pic_top = Inches(2.0) + (box_h - pic_h) / 2
            else:
                pic_h = box_h
                pic_w = box_h * aspect
                pic_top = Inches(2.0)
                pic_left = Inches(8.15) + (box_w - pic_w) / 2
                
            slide.shapes.add_picture(img_path, pic_left, pic_top, width=pic_w, height=pic_h)
            
            # Watermark overlay for photoshoot photos
            is_info = img_name.startswith("slide") or any(k in img_name for k in ["_protocol", "_card", "_station", "_care", "_admin", "_kit", "_allergies", "_positions", "_first_aid", "_dehydration", "_hd.png"])
            if not is_info:
                wm_path = os.path.join(ASSETS_DIR, "logo_horizontal_light.png")
                if os.path.exists(wm_path):
                    slide.shapes.add_picture(wm_path, Inches(9.3), Inches(2.05), width=Inches(1.8))

    # Provera za metrics traku
    metrics = s.get("metrics", [])
    top_pos = Inches(1.85)
    total_avail = 5.0
    
    if metrics:
        m_count = len(metrics)
        m_gap = 0.14
        total_w = 6.9
        m_w = (total_w - (m_count - 1) * m_gap) / m_count
        m_h = 0.98
        left_pos = Inches(0.8)
        
        for m in metrics:
            m_card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, left_pos, top_pos, Inches(m_w), Inches(m_h))
            m_card.fill.solid()
            m_card.fill.fore_color.rgb = RGB_DARK_CARD if is_dark else RGB_CARD_BG
            m_card.line.color.rgb = RGB_SAND if is_dark else RGB_BORDER
            m_card.line.width = Pt(1)
            
            bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, left_pos, top_pos, Inches(0.06), Inches(m_h))
            bar.fill.solid()
            bar.fill.fore_color.rgb = RGB_GOLD if is_dark else RGB_TERRA
            bar.line.fill.background()
            
            mtf = m_card.text_frame
            mtf.word_wrap = True
            mtf.margin_left = Inches(0.16)
            mtf.margin_right = Inches(0.1)
            mtf.margin_top = Inches(0.12)
            mtf.margin_bottom = Inches(0.08)
            
            mp1 = mtf.paragraphs[0]
            mp1.text = m["val"]
            mp1.font.name = FONT_TITLE
            mp1.font.size = Pt(21)
            mp1.font.bold = True
            mp1.font.color.rgb = RGB_GOLD if is_dark else RGB_TERRA
            
            mp2 = mtf.add_paragraph()
            mp2.text = m["lbl"]
            mp2.font.name = FONT_BODY
            mp2.font.size = Pt(10)
            mp2.font.color.rgb = RGB_LINEN if is_dark else RGB_OLIVE
            mp2.space_before = Pt(2)
            
            left_pos += Inches(m_w + m_gap)
            
        top_pos += Inches(m_h + 0.14)
        total_avail -= (m_h + 0.14)

    # Leve kartice sa tekstom - popunjavaju preostali vertikalni prostor
    blocks = s.get("content_blocks", [])
    b_count = max(1, len(blocks))
    gap = 0.14
    card_h = Inches((total_avail - (b_count - 1) * gap) / b_count)
    
    total_bullets = sum(len(b.get("bullets", [])) for b in blocks)
    is_dense = (bool(metrics) and total_bullets >= 6) or total_bullets >= 8
    
    # Dinamičko skaliranje fonta prema gustini i broju kartica
    if is_dense:
        f_title = Pt(18)
        f_desc = Pt(14)
        f_bullet = Pt(13)
        space_b = Pt(3.5)
    elif b_count == 1:
        f_title = Pt(23.5)
        f_desc = Pt(16)
        f_bullet = Pt(14.5)
        space_b = Pt(6)
    elif b_count == 2:
        f_title = Pt(20)
        f_desc = Pt(14.5)
        f_bullet = Pt(13.5)
        space_b = Pt(5)
    else:
        f_title = Pt(17.5)
        f_desc = Pt(13)
        f_bullet = Pt(12.2)
        space_b = Pt(3)
    
    for b in blocks:
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, Inches(0.8), top_pos, Inches(6.9), card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = RGB_DARK_CARD if is_dark else RGB_CARD_BG
        card.line.color.rgb = RGB_SAND if is_dark else RGB_BORDER
        card.line.width = Pt(1)
        
        # Left luxury accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(0.8), top_pos, Inches(0.07), card_h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGB_GOLD if is_dark else RGB_TERRA
        bar.line.fill.background()
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.3)
        tf.margin_right = Inches(0.24)
        tf.margin_top = Inches(0.18)
        tf.margin_bottom = Inches(0.15)
        
        p0 = tf.paragraphs[0]
        p0.text = b.get("title", "")
        p0.font.name = FONT_TITLE
        p0.font.size = f_title
        p0.font.bold = True
        p0.font.color.rgb = RGB_GOLD if is_dark else RGB_TERRA
        
        if b.get("text"):
            p_txt = tf.add_paragraph()
            p_txt.text = clean_html(b["text"])
            p_txt.font.name = FONT_BODY
            p_txt.font.size = f_desc
            p_txt.font.color.rgb = RGB_LINEN if is_dark else RGB_DARK
            p_txt.space_before = Pt(4)
            
        if b.get("bullets"):
            for bullet in b["bullets"]:
                pb = tf.add_paragraph()
                pb.text = "• " + clean_html(bullet)
                pb.font.name = FONT_BODY
                pb.font.size = f_bullet
                pb.font.color.rgb = RGB_LINEN if is_dark else RGB_DARK
                pb.space_before = space_b
                
        top_pos += card_h + Inches(gap)

def render_two_col_layout(slide, s, is_dark):
    blocks = s.get("content_blocks", [])
    lefts = [Inches(0.8), Inches(6.8)]
    col_w = Inches(5.75)
    card_h = Inches(5.0)
    top_pos = Inches(1.85)
    
    for i, b in enumerate(blocks[:2]):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, lefts[i], top_pos, col_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = RGB_DARK_CARD if is_dark else RGB_CARD_BG
        card.line.color.rgb = RGB_SAND if is_dark else RGB_BORDER
        card.line.width = Pt(1)
        
        # Left luxury accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, lefts[i], top_pos, Inches(0.08), card_h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGB_GOLD if is_dark else RGB_TERRA
        bar.line.fill.background()
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.38)
        tf.margin_right = Inches(0.32)
        tf.margin_top = Inches(0.28)
        tf.margin_bottom = Inches(0.28)
        
        p0 = tf.paragraphs[0]
        p0.text = b.get("title", "")
        p0.font.name = FONT_TITLE
        p0.font.size = Pt(21)
        p0.font.bold = True
        p0.font.color.rgb = RGB_GOLD if is_dark else RGB_TERRA
        
        if b.get("text"):
            pt = tf.add_paragraph()
            pt.text = clean_html(b["text"])
            pt.font.name = FONT_BODY
            pt.font.size = Pt(15)
            pt.font.color.rgb = RGB_LINEN if is_dark else RGB_DARK
            pt.space_before = Pt(8)
            
        if b.get("bullets"):
            for bullet in b["bullets"]:
                pb = tf.add_paragraph()
                pb.text = "• " + clean_html(bullet)
                pb.font.name = FONT_BODY
                pb.font.size = Pt(14)
                pb.font.color.rgb = RGB_LINEN if is_dark else RGB_DARK
                pb.space_before = Pt(6)

def render_three_cards_layout(slide, s, is_dark):
    blocks = s.get("content_blocks", [])
    lefts = [Inches(0.8), Inches(4.75), Inches(8.7)]
    card_w = Inches(3.8)
    card_h = Inches(5.0)
    top_pos = Inches(1.85)
    
    for i, b in enumerate(blocks[:3]):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, lefts[i], top_pos, card_w, card_h)
        card.fill.solid()
        card.fill.fore_color.rgb = RGB_DARK_CARD if is_dark else RGB_CARD_BG
        card.line.color.rgb = RGB_SAND if is_dark else RGB_BORDER
        card.line.width = Pt(1)
        
        # Left luxury accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, lefts[i], top_pos, Inches(0.07), card_h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGB_GOLD if is_dark else RGB_TERRA
        bar.line.fill.background()
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.32)
        tf.margin_right = Inches(0.28)
        tf.margin_top = Inches(0.28)
        tf.margin_bottom = Inches(0.28)
        
        p0 = tf.paragraphs[0]
        p0.text = b.get("title", "")
        p0.font.name = FONT_TITLE
        p0.font.size = Pt(19)
        p0.font.bold = True
        p0.font.color.rgb = RGB_GOLD if is_dark else RGB_TERRA
        
        if b.get("text"):
            pt = tf.add_paragraph()
            pt.text = clean_html(b["text"])
            pt.font.name = FONT_BODY
            pt.font.size = Pt(14)
            pt.font.color.rgb = RGB_LINEN if is_dark else RGB_DARK
            pt.space_before = Pt(8)
            
        if b.get("bullets"):
            for bullet in b["bullets"]:
                pb = tf.add_paragraph()
                pb.text = "• " + clean_html(bullet)
                pb.font.name = FONT_BODY
                pb.font.size = Pt(13)
                pb.font.color.rgb = RGB_LINEN if is_dark else RGB_DARK
                pb.space_before = Pt(5)

def render_matrix_layout(slide, s, is_dark):
    blocks = s.get("content_blocks", [])
    coords = [
        (Inches(0.8), Inches(1.85)),
        (Inches(6.8), Inches(1.85)),
        (Inches(0.8), Inches(4.45)),
        (Inches(6.8), Inches(4.45)),
    ]
    box_w = Inches(5.75)
    box_h = Inches(2.4)
    
    for i, b in enumerate(blocks[:4]):
        card = slide.shapes.add_shape(MSO_SHAPE.ROUNDED_RECTANGLE, coords[i][0], coords[i][1], box_w, box_h)
        card.fill.solid()
        card.fill.fore_color.rgb = RGB_DARK_CARD if is_dark else RGB_CARD_BG
        card.line.color.rgb = RGB_SAND if is_dark else RGB_BORDER
        card.line.width = Pt(1)
        
        # Left luxury accent bar
        bar = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, coords[i][0], coords[i][1], Inches(0.07), box_h)
        bar.fill.solid()
        bar.fill.fore_color.rgb = RGB_GOLD if is_dark else RGB_TERRA
        bar.line.fill.background()
        
        tf = card.text_frame
        tf.word_wrap = True
        tf.margin_left = Inches(0.32)
        tf.margin_right = Inches(0.28)
        tf.margin_top = Inches(0.22)
        tf.margin_bottom = Inches(0.22)
        
        p0 = tf.paragraphs[0]
        p0.text = b.get("title", "")
        p0.font.name = FONT_TITLE
        p0.font.size = Pt(17.5)
        p0.font.bold = True
        p0.font.color.rgb = RGB_GOLD if is_dark else RGB_TERRA
        
        if b.get("text"):
            pt = tf.add_paragraph()
            pt.text = clean_html(b["text"])
            pt.font.name = FONT_BODY
            pt.font.size = Pt(13.5)
            pt.font.color.rgb = RGB_LINEN if is_dark else RGB_DARK
            pt.space_before = Pt(5)
            
        if b.get("bullets"):
            for bullet in b["bullets"]:
                pb = tf.add_paragraph()
                pb.text = "• " + clean_html(bullet)
                pb.font.name = FONT_BODY
                pb.font.size = Pt(12.5)
                pb.font.color.rgb = RGB_LINEN if is_dark else RGB_DARK
                pb.space_before = Pt(4)

if __name__ == "__main__":
    build_pptx()
